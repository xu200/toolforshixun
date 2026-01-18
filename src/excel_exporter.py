"""
Excel 导出模块
生成成绩汇总表
"""
import os
from typing import List
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

from .doc_parser import StudentInfo


class ExcelExporter:
    """Excel 导出器"""
    
    # 表头配置
    HEADERS = ["学号", "姓名", "文件名", "分数", "评语"]
    
    # 列宽配置
    COLUMN_WIDTHS = {
        "A": 15,   # 学号
        "B": 12,   # 姓名
        "C": 50,   # 文件名
        "D": 10,   # 分数
        "E": 80,   # 评语
    }
    
    def __init__(self):
        self.workbook = None
        self.worksheet = None
    
    def export(
        self, 
        students: List[StudentInfo], 
        output_path: str,
        course_name: str = "",
        training_name: str = ""
    ) -> tuple[bool, str]:
        """
        导出成绩汇总 Excel
        
        Args:
            students: 学生信息列表
            output_path: 输出文件路径
            course_name: 课程名称
            training_name: 实训名称
            
        Returns:
            (是否成功, 消息)
        """
        try:
            self.workbook = Workbook()
            self.worksheet = self.workbook.active
            
            # 设置工作表标题
            sheet_title = training_name if training_name else "成绩汇总"
            self.worksheet.title = sheet_title[:31]  # Excel 限制 31 字符
            
            # 写入表头
            self._write_header()
            
            # 写入数据
            self._write_data(students)
            
            # 应用样式
            self._apply_styles(len(students))
            
            # 保存文件
            self.workbook.save(output_path)
            
            return True, f"成功导出 {len(students)} 条记录到 {os.path.basename(output_path)}"
            
        except Exception as e:
            return False, f"导出失败: {str(e)}"
        
        finally:
            if self.workbook:
                self.workbook.close()
    
    def _write_header(self):
        """写入表头"""
        for col, header in enumerate(self.HEADERS, start=1):
            cell = self.worksheet.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, name="宋体", size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center")
    
    def _write_data(self, students: List[StudentInfo]):
        """写入学生数据"""
        for row, student in enumerate(students, start=2):
            self.worksheet.cell(row=row, column=1, value=student.student_id)
            self.worksheet.cell(row=row, column=2, value=student.student_name)
            self.worksheet.cell(row=row, column=3, value=student.file_name)
            self.worksheet.cell(row=row, column=4, value=student.score if student.score is not None else "")
            self.worksheet.cell(row=row, column=5, value=student.comment)
    
    def _apply_styles(self, data_rows: int):
        """应用样式"""
        # 设置列宽
        for col, width in self.COLUMN_WIDTHS.items():
            self.worksheet.column_dimensions[col].width = width
        
        # 设置边框
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # 应用边框和对齐
        for row in range(1, data_rows + 2):
            for col in range(1, len(self.HEADERS) + 1):
                cell = self.worksheet.cell(row=row, column=col)
                cell.border = thin_border
                cell.font = Font(name="宋体", size=10)
                
                # 评语列左对齐，其他居中
                if col == 5:
                    cell.alignment = Alignment(
                        horizontal="left", 
                        vertical="center",
                        wrap_text=True
                    )
                else:
                    cell.alignment = Alignment(
                        horizontal="center", 
                        vertical="center"
                    )
        
        # 设置行高
        for row in range(2, data_rows + 2):
            self.worksheet.row_dimensions[row].height = 60  # 评语可能较长
    
    def generate_filename(self, course_name: str, training_name: str) -> str:
        """
        生成导出文件名
        
        Args:
            course_name: 课程名称
            training_name: 实训名称
            
        Returns:
            文件名
        """
        if course_name and training_name:
            return f"{course_name}_{training_name}_成绩汇总.xlsx"
        elif training_name:
            return f"{training_name}_成绩汇总.xlsx"
        else:
            return "实训成绩汇总.xlsx"
