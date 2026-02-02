"""
标红模块
负责将不一致的单元格标红并保存为新文件
"""
import pandas as pd
from typing import List, Tuple, Any
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows


class ScoreHighlighter:
    """成绩标红器"""
    
    RED_FONT = Font(color="FF0000", bold=True)
    RED_FILL = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
    
    def __init__(self):
        pass
    
    def highlight_and_save(
        self,
        manual_df: pd.DataFrame,
        mismatches: List[Tuple[int, str, Any, Any]],
        output_path: str
    ) -> Tuple[bool, str]:
        """
        将不一致的单元格标红并保存为新文件
        
        Args:
            manual_df: 人工 Excel 的 DataFrame
            mismatches: 不一致列表 [(行索引, 列名, 人工分数, 系统分数), ...]
            output_path: 输出文件路径
            
        Returns:
            (是否成功, 消息)
        """
        if manual_df is None or manual_df.empty:
            return False, "没有数据可导出"
        
        try:
            manual_df.to_excel(output_path, index=False, sheet_name="对比结果")
            
            wb = load_workbook(output_path)
            ws = wb.active
            
            columns = list(manual_df.columns)
            
            mismatch_positions = set()
            for row_idx, col_name, manual_score, auto_score in mismatches:
                if col_name in columns:
                    col_idx = columns.index(col_name) + 1
                    excel_row = row_idx + 2
                    mismatch_positions.add((excel_row, col_idx))
            
            for excel_row, col_idx in mismatch_positions:
                cell = ws.cell(row=excel_row, column=col_idx)
                cell.font = self.RED_FONT
                cell.fill = self.RED_FILL
            
            wb.save(output_path)
            wb.close()
            
            return True, f"成功导出对比结果，共标红 {len(mismatch_positions)} 处"
            
        except PermissionError:
            return False, "文件被占用，请关闭后重试"
        except Exception as e:
            return False, f"导出失败: {str(e)}"
