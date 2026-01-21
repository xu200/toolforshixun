"""
Word 文档解析模块
从实训报告中提取学生信息
"""
import os
import io
import re
from dataclasses import dataclass, field
from typing import List, Optional, Tuple
from docx import Document
from src.lab_name_fixer import LabNameFixer


@dataclass
class StudentInfo:
    """学生信息数据类"""
    file_path: str = ""
    file_name: str = ""
    student_id: str = ""
    student_name: str = ""
    course_name: str = ""
    training_name: str = ""
    lab_name: str = ""
    score: Optional[float] = None
    score_source: str = ""
    word_count: int = 0
    image_count: int = 0
    comment: str = ""
    is_valid: bool = True
    error_message: str = ""
    
    def get_new_filename(self) -> str:
        """生成规范化文件名（限制长度避免路径过长）"""
        if not all([self.student_id, self.student_name, self.course_name, self.training_name]):
            return self.file_name
        
        # 截取各字段长度，避免文件名过长
        course = self.course_name[:15] if len(self.course_name) > 15 else self.course_name
        training = self.training_name[:20] if len(self.training_name) > 20 else self.training_name
        
        # 清理特殊字符
        course = re.sub(r'[\\/:*?"<>|]', '', course)
        training = re.sub(r'[\\/:*?"<>|]', '', training)
        
        return f"{self.student_id}_{self.student_name}_{course}_{training}.docx"


class DocParser:
    """Word 文档解析器"""
    
    def __init__(self):
        self.students: List[StudentInfo] = []
        self.course_name: str = ""
        self.training_name: str = ""
        self.min_words: int = 0
        self.max_words: int = 0
        self.lab_name_fixer = LabNameFixer()
    
    # 字段匹配模式（增强版，支持多种格式）
    PATTERNS = {
        "student_name": [
            r"学生姓名[：:\s]*([^\s\n\r：:]+)",
            r"姓\s*名[：:\s]*([^\s\n\r：:]+)",
            r"学\s*生[：:\s]*([^\s\n\r：:]+)",
            r"姓名[：:\s]+(\S+)",
            r"姓\s*名\s*[：:]\s*([\u4e00-\u9fa5]{2,4})",  # 支持表格中的中文姓名
        ],
        "student_id": [
            r"学\s*号[：:\s]*(\d{6,12})",
            r"学生学号[：:\s]*(\d{6,12})",
            r"学号[：:\s]+(\d+)",
            r"学\s*号\s*[：:]\s*(\d{6,12})",  # 支持表格中带空格的格式
            r"(\d{9,12})",  # 直接匹配9-12位数字作为学号
        ],
        "course_name": [
            r"课程名称[：:\s]*([^\n\r]+?)(?=\s{2,}|\n|$)",
            r"课\s*程[：:\s]*([^\n\r]+?)(?=\s{2,}|\n|$)",
            r"课程[：:\s]+([^\n\r]+?)(?=\s{2,}|\n|$)",
        ],
        "training_name": [
            r"实验[（\(]训[）\)]名称[：:\s]*([^\n\r，,]{2,30})",  # 匹配"实验（训）名称"
            r"实训名称[：:\s]*([^\n\r，,]{2,30})",
            r"实训项目[：:\s]*([^\n\r，,]{2,30})",
            r"实训主题[：:\s]*([^\n\r，,]{2,30})",
            r"项目名称[：:\s]*([^\n\r，,]{2,30})",
        ],
    }

    def parse_folder(self, folder_path: str) -> List[StudentInfo]:
        """
        解析文件夹中的所有 Word 文档
        
        Args:
            folder_path: 文件夹路径
            
        Returns:
            学生信息列表
        """
        self.students = []
        
        if not os.path.isdir(folder_path):
            return self.students
        
        for filename in os.listdir(folder_path):
            if filename.endswith('.docx') and not filename.startswith('~$'):
                file_path = os.path.join(folder_path, filename)
                student_info = self.parse_document(file_path)
                self.students.append(student_info)

        word_counts = [s.word_count for s in self.students if s.word_count > 0]
        if word_counts:
            self.min_words = min(word_counts)
            self.max_words = max(word_counts)
        else:
            self.min_words = 0
            self.max_words = 0

        return self.students

    def apply_auto_scores(self, students: List[StudentInfo]) -> int:
        """手动触发：根据字数/图片数建议分数（不覆盖教师评分/手动评分）"""
        word_counts = [s.word_count for s in students if s.word_count > 0 and s.is_valid]
        if word_counts:
            self.min_words = min(word_counts)
            self.max_words = max(word_counts)
        else:
            self.min_words = 0
            self.max_words = 0

        changed = 0
        for student in students:
            if not student.is_valid:
                continue
            if getattr(student, "score_source", "") == "teacher":
                continue
            if getattr(student, "score_source", "") == "manual":
                continue

            suggested = self.suggest_score(student.word_count, student.image_count)
            if suggested is None:
                continue

            student.score = float(int(suggested))
            student.score_source = "auto"
            changed += 1

        return changed
    
    def parse_document(self, file_path: str) -> StudentInfo:
        """
        解析单个 Word 文档
        
        Args:
            file_path: 文件路径
            
        Returns:
            学生信息
        """
        student = StudentInfo(
            file_path=file_path,
            file_name=os.path.basename(file_path)
        )
        
        try:
            with open(file_path, "rb") as f:
                data = f.read()
            doc = Document(io.BytesIO(data))
            full_text = self._extract_text(doc)
            
            # 提取各字段
            student.student_name = self._extract_field(full_text, "student_name")
            student.student_id = self._extract_field(full_text, "student_id")
            student.course_name = self._extract_field(full_text, "course_name")
            student.training_name = self._extract_field(full_text, "training_name")
            
            # 如果文档内解析失败，尝试从文件名回退提取
            if not student.student_id or not student.student_name:
                fallback_id, fallback_name = self._extract_from_filename(student.file_name)
                if not student.student_id and fallback_id:
                    student.student_id = fallback_id
                if not student.student_name and fallback_name:
                    student.student_name = fallback_name
            
            try:
                student.lab_name = self.lab_name_fixer.extract_lab_name_from_doc(doc)
            except Exception as e:
                student.lab_name = ""

            teacher_score = self._extract_teacher_score(doc)
            if teacher_score is not None:
                student.score = float(int(teacher_score))
                student.score_source = "teacher"

            student.word_count = self._count_effective_words(full_text)
            student.image_count = self._count_images(doc)
            
            # 检查必要字段（放宽条件：只要有字数统计就认为可用）
            missing_fields = []
            if not student.student_id:
                missing_fields.append("学号")
            if not student.student_name:
                missing_fields.append("姓名")
            
            if missing_fields:
                student.error_message = f"缺少字段: {', '.join(missing_fields)}"
                # 放宽条件：只要有字数统计就认为可用（可以参与自动评分）
                if student.word_count > 0:
                    student.is_valid = True  # 仍然可用，只是有警告
                else:
                    student.is_valid = False
                
        except Exception as e:
            student.is_valid = False
            student.error_message = f"解析错误: {str(e)}"
        
        return student

    def suggest_score(self, word_count: int, image_count: int) -> Optional[int]:
        if word_count <= 0:
            return None
        if self.min_words <= 0 and self.max_words <= 0:
            return None

        word_range = self.max_words - self.min_words
        interval_size = word_range / 10 if word_range > 0 else 0
        if interval_size <= 0:
            index = 0
        else:
            index = int((word_count - self.min_words) // interval_size)
        if index < 0:
            index = 0
        if index > 10:
            index = 10

        text_score = 80 + index
        if text_score < 80:
            text_score = 80
        if text_score > 90:
            text_score = 90

        if text_score == 90:
            bonus = image_count if image_count > 0 else 0
            if bonus > 5:
                bonus = 5
            final_score = 90 + bonus
            if final_score > 95:
                final_score = 95
            return int(final_score)
        return int(text_score)
    
    def _extract_text(self, doc: Document) -> str:
        """提取文档全部文本"""
        paragraphs = []
        for para in doc.paragraphs:
            paragraphs.append(para.text)
        
        # 同时提取表格中的文本
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.append(cell.text)
        
        return '\n'.join(paragraphs)

    def _count_effective_words(self, full_text: str) -> int:
        text = full_text
        parts = re.split(r"教师评阅|教师评语", text, maxsplit=1)
        if parts:
            text = parts[0]
        chars = re.findall(r"[\u4e00-\u9fffA-Za-z0-9]", text)
        return len(chars)

    def _count_images(self, doc: Document) -> int:
        try:
            return len(doc.inline_shapes)
        except Exception:
            return 0

    def _extract_teacher_score(self, doc: Document) -> Optional[int]:
        def iter_paragraph_sequences():
            yield doc.paragraphs
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield cell.paragraphs

        def parse_score(text: str) -> Optional[int]:
            cleaned_text = text
            cleaned_text = re.sub(
                r"(19|20)\d{2}\s*[年\-/\.]\s*\d{1,2}\s*[月\-/\.]\s*\d{1,2}\s*日?",
                " ",
                cleaned_text,
            )
            cleaned_text = re.sub(r"(19|20)\d{2}\s+\d{1,2}\s+\d{1,2}", " ", cleaned_text)

            for m in re.finditer(r"(?<!\d)(\d{1,3})(?!\d)", cleaned_text):
                v = int(m.group(1))
                if v < 0 or v > 100:
                    continue
                window = cleaned_text[max(0, m.start() - 2): min(len(cleaned_text), m.end() + 2)]
                if any(ch in window for ch in ("年", "月", "日")):
                    continue
                if any(ch in window for ch in ("-", "/", ".")):
                    continue
                return v
            return None

        stop_tokens = ("教师签名", "签名", "日期")
        hit_tokens = ("教师评阅", "教师评语")

        for seq in iter_paragraph_sequences():
            for idx, para in enumerate(seq):
                t = (para.text or "").strip()
                if not t:
                    continue
                if any(k in t for k in hit_tokens):
                    parts: List[str] = []

                    t2 = t
                    pos = t2.find("：")
                    if pos < 0:
                        pos = t2.find(":")
                    if pos >= 0 and pos + 1 < len(t2):
                        parts.append(t2[pos + 1 :].strip())
                    else:
                        cleaned = re.sub(r".*?(教师评阅|教师评语)\s*[:：]?", "", t2).strip()
                        if cleaned:
                            parts.append(cleaned)

                    for j in range(idx + 1, min(idx + 3, len(seq))):
                        tj = (seq[j].text or "").strip()
                        if not tj:
                            continue
                        if any(s in tj for s in stop_tokens):
                            break
                        parts.append(tj)

                    candidate = "\n".join([p for p in parts if p])
                    score = parse_score(candidate)
                    return score
        return None
    
    def _extract_field(self, text: str, field_name: str) -> str:
        """
        从文本中提取指定字段
        
        Args:
            text: 文档文本
            field_name: 字段名
            
        Returns:
            提取的值
        """
        patterns = self.PATTERNS.get(field_name, [])
        
        for pattern in patterns:
            match = re.search(pattern, text)
            if match:
                value = match.group(1).strip()
                # 清理提取的值
                value = re.sub(r'\s+', '', value)
                return value
        
        return ""
    
    def _extract_from_filename(self, filename: str) -> Tuple[str, str]:
        """
        从文件名中提取学号和姓名（作为回退方案）
        
        支持的文件名格式：
        - 202303716-B230602-赖宏-第一次实训.docx
        - 202303720-B230602-杨文博-第一次实训.docx
        - 学号_姓名_课程_实训.docx
        
        Args:
            filename: 文件名
            
        Returns:
            (学号, 姓名)
        """
        # 去掉扩展名
        name = os.path.splitext(filename)[0]
        
        # 尝试多种分隔符
        for sep in ['-', '_', ' ']:
            parts = name.split(sep)
            if len(parts) >= 2:
                # 查找学号（纯数字，6-12位）
                student_id = ""
                student_name = ""
                
                for i, part in enumerate(parts):
                    part = part.strip()
                    # 学号：纯数字6-12位
                    if re.match(r'^\d{6,12}$', part) and not student_id:
                        student_id = part
                    # 姓名：2-4个中文字符
                    elif re.match(r'^[\u4e00-\u9fa5]{2,4}$', part) and not student_name:
                        student_name = part
                
                if student_id or student_name:
                    return student_id, student_name
        
        return "", ""
    
    def get_course_info(self) -> Tuple[str, str]:
        """
        获取课程信息（从第一个有效文档）
        
        Returns:
            (课程名称, 实训名称)
        """
        for student in self.students:
            if student.course_name and student.training_name:
                return student.course_name, student.training_name
        return "", ""
