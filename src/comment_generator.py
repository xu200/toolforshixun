"""
评语生成与写回模块
生成评语并写入 Word 文档
"""
import io
import os
import stat
import time
from typing import List, Dict, Set, Tuple, Optional
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

from .comment_rules import generate_comment, get_level_by_score
from .doc_parser import StudentInfo


class CommentGenerator:
    """评语生成器"""
    
    def __init__(self):
        # 按等级记录已使用的评语组合，避免重复
        self.used_combinations: Dict[str, Set[Tuple[int, int, int]]] = {}
    
    def reset_combinations(self):
        """重置已使用的组合记录"""
        self.used_combinations = {}
    
    def generate_for_student(self, student: StudentInfo) -> str:
        """
        为单个学生生成评语
        
        Args:
            student: 学生信息
            
        Returns:
            生成的评语
        """
        if student.score is None:
            return ""
        
        level = get_level_by_score(student.score)
        
        # 获取该等级已使用的组合
        if level not in self.used_combinations:
            self.used_combinations[level] = set()
        
        comment, combination = generate_comment(
            student.score, 
            self.used_combinations[level]
        )
        
        # 记录使用的组合
        self.used_combinations[level].add(combination)
        
        student.comment = comment
        return comment
    
    def generate_for_all(self, students: List[StudentInfo]) -> int:
        """
        为所有学生批量生成评语
        
        Args:
            students: 学生信息列表
            
        Returns:
            成功生成评语的数量
        """
        self.reset_combinations()
        count = 0
        
        for student in students:
            if student.score is not None:
                self.generate_for_student(student)
                count += 1
        
        return count
    
    def regenerate_for_student(self, student: StudentInfo) -> str:
        """
        为单个学生重新生成评语
        
        Args:
            student: 学生信息
            
        Returns:
            新生成的评语
        """
        if student.score is None:
            return ""
        
        level = get_level_by_score(student.score)
        
        # 不使用已记录的组合，强制生成新的
        comment, _ = generate_comment(student.score)
        student.comment = comment
        return comment


class CommentWriter:
    """评语写入器"""
    
    # 写入格式配置
    FONT_NAME = "宋体"
    FONT_SIZE = Pt(10.5)  # 五号字
    
    def write_to_document(self, student: StudentInfo) -> Tuple[bool, str]:
        """
        将评语写入 Word 文档
        
        Args:
            student: 学生信息
            
        Returns:
            (是否成功, 消息)
        """
        if not student.comment:
            return False, "评语为空"
        
        file_path = student.file_path
        if not file_path:
            return False, "文件路径为空"

        file_path = os.path.normpath(file_path)
        student.file_path = file_path

        if not os.path.exists(file_path):
            return False, f"文件不存在: {file_path}"
        
        try:
            with open(file_path, "rb") as f:
                data = f.read()
            doc = Document(io.BytesIO(data))
            
            # 查找"教师评阅"位置并写入
            success = self._find_and_write(doc, student.comment)
            
            if not success:
                return False, f"未找到教师评阅位置 | 路径: {file_path}"
            
            tmp_path = f"{file_path}.tmp"
            last_error: Optional[Exception] = None
            for attempt in range(3):
                try:
                    doc.save(tmp_path)
                    try:
                        os.chmod(file_path, stat.S_IWRITE)
                    except Exception:
                        pass
                    os.replace(tmp_path, file_path)
                    return True, "写入成功"
                except Exception as e:
                    last_error = e
                    winerror = getattr(e, 'winerror', None)
                    if attempt >= 2 and winerror == 32:
                        try:
                            root, ext = os.path.splitext(file_path)
                            fallback_path = f"{root}_待替换{ext}"
                            idx = 1
                            while os.path.exists(fallback_path):
                                fallback_path = f"{root}_待替换{idx}{ext}"
                                idx += 1
                            os.replace(tmp_path, fallback_path)
                            return False, (
                                f"原文件被占用(WinError 32)，无法覆盖保存: {file_path}"
                                f"\n已生成待替换文件(评语已写入): {fallback_path}"
                                f"\n请关闭占用程序后再次写入，或手动用待替换文件替换原文件"
                            )
                        except Exception:
                            pass
                    try:
                        if os.path.exists(tmp_path):
                            os.remove(tmp_path)
                    except Exception:
                        pass
                    if attempt < 2:
                        time.sleep(0.4 * (attempt + 1))
                        continue

            winerror = getattr(last_error, 'winerror', None) if last_error else None
            if winerror == 32:
                return False, (
                    f"文件被占用(WinError 32): {file_path}"
                    f"\n可能原因: Word/资源管理器预览窗格/OneDrive同步/杀毒软件正在占用"
                )
            if last_error:
                return False, f"写入失败: {type(last_error).__name__}: {last_error} | 路径: {file_path}"
            return False, f"写入失败: 未知错误 | 路径: {file_path}"
        except Exception as e:
            import traceback
            err_type = type(e).__name__
            winerror = getattr(e, 'winerror', None)
            if winerror is not None:
                brief = f"{err_type}(WinError {winerror}): {e}"
            else:
                brief = f"{err_type}: {e}"
            error_detail = traceback.format_exc()
            print(f"写入错误详情(路径={file_path}):\n{error_detail}")
            return False, f"{brief} | 路径: {file_path}"
    
    def _find_and_write(self, doc: Document, comment: str) -> bool:
        """
        查找教师评阅位置并写入评语（支持段落和表格）
        
        Args:
            doc: Word 文档对象
            comment: 评语内容
            
        Returns:
            是否找到并写入
        """
        # 先在段落中查找
        for i, para in enumerate(doc.paragraphs):
            if "教师评阅" in para.text or "教师评语" in para.text:
                # 在下一段落写入，或者在当前段落后添加
                if "：" in para.text or ":" in para.text:
                    idx = para.text.find("：")
                    if idx < 0:
                        idx = para.text.find(":")
                    prefix = para.text[: idx + 1] if idx >= 0 else "教师评阅："
                    para.clear()
                    title_run = para.add_run(prefix)
                    self._apply_format(title_run)
                    comment_run = para.add_run(comment)
                    self._apply_format(comment_run)
                elif i + 1 < len(doc.paragraphs):
                    next_para = doc.paragraphs[i + 1]
                    # 清空原有内容并写入新评语
                    next_para.clear()
                    run = next_para.add_run(comment)
                    self._apply_format(run)
                else:
                    # 添加新段落
                    new_para = doc.add_paragraph()
                    run = new_para.add_run(comment)
                    self._apply_format(run)
                return True
        
        # 在表格中查找（关键修复：很多实训报告的教师评阅在表格中）
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    # 遍历单元格中的每个段落
                    for para in cell.paragraphs:
                        if "教师评阅" in para.text:
                            # 找到"教师评阅："所在段落，直接在后面追加评语
                            # 不清空，不新建段落，直接追加文字
                            comment_run = para.add_run(comment)
                            self._apply_format(comment_run)
                            return True
        
        return False
    
    def _append_comment(self, doc: Document, comment: str):
        """
        在文档末尾添加评语
        
        Args:
            doc: Word 文档对象
            comment: 评语内容
        """
        # 添加教师评阅标题
        title_para = doc.add_paragraph()
        title_run = title_para.add_run("教师评阅：")
        self._apply_format(title_run)
        
        # 添加评语内容
        comment_para = doc.add_paragraph()
        comment_run = comment_para.add_run(comment)
        self._apply_format(comment_run)
    
    def _apply_format(self, run):
        """
        应用格式设置
        
        Args:
            run: Word run 对象
        """
        run.font.name = self.FONT_NAME
        run.font.size = self.FONT_SIZE
        run.font.bold = False
        run.font.italic = False
        
        # 安全设置中文字体
        r = run._element
        rPr = r.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn('w:eastAsia'), self.FONT_NAME)
    
    def write_all(self, students: List[StudentInfo]) -> List[Tuple[str, bool, str]]:
        """
        批量写入评语到所有文档
        
        Args:
            students: 学生信息列表
            
        Returns:
            写入结果列表: [(文件名, 是否成功, 消息)]
        """
        results = []
        
        for student in students:
            if student.comment:
                success, message = self.write_to_document(student)
                results.append((student.file_name, success, message))
            else:
                results.append((student.file_name, False, "无评语"))
        
        return results
